from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_AUTO_SIZE
from pptx.util import Inches, Pt


OUT = Path("AiQuezBot_parallel_qismlar.pptx")


def set_run(run, *, size=20, bold=False, color=(30, 41, 59), font="Aptos"):
    run.font.name = font
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = RGBColor(*color)


def add_textbox(slide, x, y, w, h, text="", *, size=20, bold=False, color=(30, 41, 59)):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.auto_size = MSO_AUTO_SIZE.TEXT_TO_FIT_SHAPE
    p = tf.paragraphs[0]
    p.text = text
    p.alignment = PP_ALIGN.LEFT
    for run in p.runs:
        set_run(run, size=size, bold=bold, color=color)
    return box


def add_title(slide, title, subtitle=None):
    add_textbox(slide, 0.55, 0.35, 12.2, 0.55, title, size=28, bold=True, color=(15, 23, 42))
    if subtitle:
        add_textbox(slide, 0.58, 0.9, 11.8, 0.35, subtitle, size=13, color=(71, 85, 105))


def add_bullets(slide, items, x=0.72, y=1.35, w=5.7, h=4.8, size=16):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.auto_size = MSO_AUTO_SIZE.TEXT_TO_FIT_SHAPE
    for i, item in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = item
        p.level = 0
        p.space_after = Pt(8)
        for run in p.runs:
            set_run(run, size=size, color=(30, 41, 59))
    return box


def add_code(slide, code, x=6.55, y=1.35, w=6.15, h=4.75, font_size=12):
    shape = slide.shapes.add_shape(1, Inches(x), Inches(y), Inches(w), Inches(h))
    shape.fill.solid()
    shape.fill.fore_color.rgb = RGBColor(15, 23, 42)
    shape.line.color.rgb = RGBColor(51, 65, 85)
    tf = shape.text_frame
    tf.clear()
    tf.margin_left = Inches(0.22)
    tf.margin_right = Inches(0.15)
    tf.margin_top = Inches(0.16)
    tf.margin_bottom = Inches(0.12)
    tf.word_wrap = False
    p = tf.paragraphs[0]
    p.text = code
    for run in p.runs:
        set_run(run, size=font_size, color=(226, 232, 240), font="Cascadia Mono")
    return shape


def add_footer(slide, left, right="AiQuezBot"):
    add_textbox(slide, 0.55, 7.05, 8.8, 0.22, left, size=9, color=(100, 116, 139))
    box = add_textbox(slide, 10.7, 7.05, 2.1, 0.22, right, size=9, color=(100, 116, 139))
    box.text_frame.paragraphs[0].alignment = PP_ALIGN.RIGHT


def add_section_slide(prs, title, subtitle, bullets):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    bg = slide.background.fill
    bg.solid()
    bg.fore_color.rgb = RGBColor(248, 250, 252)
    add_title(slide, title, subtitle)
    add_bullets(slide, bullets, x=0.85, y=1.65, w=11.7, h=4.6, size=18)
    add_footer(slide, "Qisqa izoh: parallelizm bu yerda asosan asyncio event loop orqali ishlaydi.")
    return slide


def main():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = RGBColor(241, 245, 249)
    add_textbox(slide, 0.75, 0.8, 11.9, 1.0, "AiQuezBot'da paralellashtirish", size=36, bold=True, color=(15, 23, 42))
    add_textbox(
        slide,
        0.8,
        1.85,
        11.4,
        0.8,
        "Botdagi parallel ishlov berilgan qismlar: qayerda ishlatilgani, qaysi funksiya ekani va qisqa kod parchalari.",
        size=19,
        color=(51, 65, 85),
    )
    add_bullets(
        slide,
        [
            "Asosiy texnologiya: Python asyncio",
            "Parallel API chaqiruvlar: asyncio.gather(...)",
            "Fon jarayonlar: asyncio.create_task(...)",
            "Bloklovchi ishlarni ajratish: asyncio.to_thread(...)",
        ],
        x=0.95,
        y=3.05,
        w=11.4,
        h=2.3,
        size=20,
    )
    add_footer(slide, "Tayyorlangan fayl: AiQuezBot_parallel_qismlar.pptx")

    add_section_slide(
        prs,
        "1. Umumiy yondashuv",
        "Bot bir vaqtning o'zida foydalanuvchi xabarlari, AI chaqiruvlari, fayl ishlovlari va quiz jarayonlarini boshqaradi.",
        [
            "Kodning katta qismi async def va await bilan yozilgan.",
            "Event loop I/O kutish vaqtida boshqa vazifalarni bajarishga imkon beradi.",
            "Asosiy parallel nuqtalar: create_task, gather, wait_for, to_thread, Event, Lock.",
            "Bu usul Telegram botlarda javob berishni sekinlashtirmaslik uchun muhim.",
        ],
    )

    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = RGBColor(248, 250, 252)
    add_title(slide, "2. Fon task: keepalive", "main.py:246 va main.py:270")
    add_bullets(
        slide,
        [
            "create_task keepalive loopni asosiy polling/webhook bilan parallel ishlatadi.",
            "Bot asosiy ishini davom ettiradi, keepalive esa orqa fonda URLga so'rov yuboradi.",
            "finally blokida task cancel qilinadi, ya'ni ilova yopilganda toza tugatiladi.",
        ],
    )
    add_code(
        slide,
        """keepalive_task = asyncio.create_task(
    _keepalive_loop(url, interval_sec=interval)
)

# ...
if keepalive_task is not None:
    keepalive_task.cancel()
    try:
        await keepalive_task
    except asyncio.CancelledError:
        pass""",
        font_size=13,
    )
    add_footer(slide, "Funksiya: asyncio.create_task")

    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = RGBColor(248, 250, 252)
    add_title(slide, "3. AI savol yaratish batchlari", "services/ai_service.py:1151-1155 va 1233")
    add_bullets(
        slide,
        [
            "Matn bir nechta bo'lakka bo'linadi va har bo'lakdan savol yaratish alohida task bo'ladi.",
            "asyncio.gather barcha AI chaqiruvlarni bir vaqtda kutadi.",
            "return_exceptions=True bitta batch xato qilsa ham qolgan natijalarni saqlab qoladi.",
            "Tezlik sozlamalari: AI_BATCH_SIZE, AI_PARALLEL_BATCHES_TEXT.",
        ],
    )
    add_code(
        slide,
        """init_tasks = [
    _make_text_task(text_chunks[i % len(text_chunks)],
                    min(batch_size, question_count))
    for i in range(n_initial)
]
init_results = await asyncio.gather(
    *init_tasks,
    return_exceptions=True
)""",
        font_size=12.5,
    )
    add_footer(slide, "Funksiya: asyncio.gather")

    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = RGBColor(248, 250, 252)
    add_title(slide, "4. Yetmayotgan savollarni parallel to'ldirish", "services/ai_service.py:1200-1233")
    add_bullets(
        slide,
        [
            "Agar kerakli savollar soni to'lmasa, kod parallel batchlar bilan yana so'rov yuboradi.",
            "parallel qiymati 1-5 oralig'ida cheklanadi.",
            "Har bir task boshqa chunk bilan ishlaydi, shuning uchun takror savollar kamayadi.",
        ],
    )
    add_code(
        slide,
        """parallel = int(os.getenv(
    "AI_PARALLEL_BATCHES_TEXT",
    os.getenv("AI_PARALLEL_BATCHES", "2")
))
parallel = max(1, min(5, parallel))

tasks = []
for j in range(parallel):
    tasks.append(self._generate_gemini(chunk, ask, ...))

results = await asyncio.gather(
    *tasks,
    return_exceptions=True
)""",
        font_size=11.5,
    )
    add_footer(slide, "Funksiya: asyncio.gather + parallel batch konfiguratsiyasi")

    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = RGBColor(248, 250, 252)
    add_title(slide, "5. AI javoblarni tekshirish", "services/ai_service.py:1303-1306")
    add_bullets(
        slide,
        [
            "Savollar review_batch bo'yicha guruhlanadi.",
            "Har bir guruh alohida AI review task sifatida ishlaydi.",
            "Bu ko'p savolli quizlarda tekshirish vaqtini sezilarli qisqartiradi.",
        ],
    )
    add_code(
        slide,
        """async def _review_one(batch, rel_text):
    if provider == "openai":
        return await self._review_openai_answers(rel_text, batch)
    if provider == "gemini":
        return await self._review_gemini_answers(rel_text, batch)

all_fixes = await asyncio.gather(
    *[_review_one(b, t) for b, t in batch_pairs],
    return_exceptions=True,
)""",
        font_size=11.8,
    )
    add_footer(slide, "Funksiya: asyncio.gather")

    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = RGBColor(248, 250, 252)
    add_title(slide, "6. Topic context manbalari parallel olinadi", "services/topic_context_service.py:284-294")
    add_bullets(
        slide,
        [
            "Google Books, OpenLibrary va Wikipedia so'rovlari bir vaqtda yuboriladi.",
            "wait_for umumiy vaqt limitini qo'yadi.",
            "Natija tezroq qaytadi, chunki eng sekin manba butun jarayonni cheksiz ushlab turmaydi.",
        ],
    )
    add_code(
        slide,
        """tasks = [
    _google_books(session, raw, timeout_sec=timeout_sec),
    _openlibrary(session, raw, timeout_sec=timeout_sec),
]
for wl in wiki_langs[:3]:
    tasks.append(_wikipedia_summary(session, raw, lang=wl,
                                    timeout_sec=timeout_sec))

done = await asyncio.wait_for(
    asyncio.gather(*tasks, return_exceptions=True),
    timeout=timeout_sec + 1.5,
)""",
        font_size=10.8,
    )
    add_footer(slide, "Funksiyalar: asyncio.gather + asyncio.wait_for")

    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = RGBColor(248, 250, 252)
    add_title(slide, "7. Quiz jarayoni fon taskda yuradi", "handlers/user.py:1516, 414, 1624")
    add_bullets(
        slide,
        [
            "_run_quiz alohida task sifatida boshlanadi.",
            "advance_event orqali foydalanuvchi keyingi savolga o'tkazishi mumkin.",
            "wait_for savol vaqtini kutadi, lekin event kelsa oldinroq davom etadi.",
            "Shu paytda bot boshqa callback va xabarlarni ham qabul qila oladi.",
        ],
    )
    add_code(
        slide,
        """# QuizRun dataclass ichida
advance_event: asyncio.Event = field(default_factory=asyncio.Event)

# Quiz boshlanganda
run.task = asyncio.create_task(_run_quiz(bot, run))

# Har savolda kutish
run.advance_event.clear()
try:
    await asyncio.wait_for(run.advance_event.wait(),
                           timeout=run.open_period)
except asyncio.TimeoutError:
    pass""",
        font_size=11,
    )
    add_footer(slide, "Funksiyalar: create_task, Event, wait_for")

    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = RGBColor(248, 250, 252)
    add_title(slide, "8. Bloklovchi ishlar alohida threadda", "handlers/user.py:604, 6798+, 7542+")
    add_bullets(
        slide,
        [
            "DOCX eksport, fayldan matn ajratish va sahifa sanash CPU/disk ishlariga kiradi.",
            "asyncio.to_thread bu ishlarni alohida threadga beradi.",
            "Natijada event loop bloklanmaydi, bot boshqa foydalanuvchilarga javob berishda davom etadi.",
        ],
    )
    add_code(
        slide,
        """await asyncio.to_thread(
    export_quiz_to_docx,
    quiz,
    tmp_path,
    answer_title=answer_title,
    include_images=True,
)

use_text = await asyncio.to_thread(
    _extract_pptx_text_range,
    Path(pptx_path),
    page_from,
    page_to,
)""",
        font_size=11.3,
    )
    add_footer(slide, "Funksiya: asyncio.to_thread")

    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = RGBColor(248, 250, 252)
    add_title(slide, "9. Bir foydalanuvchi uchun race condition himoyasi", "handlers/user.py:128, 174-186")
    add_bullets(
        slide,
        [
            "Manual correct jarayonida bir foydalanuvchiga bitta Lock ajratiladi.",
            "Bir vaqtda kelgan callback/xabarlar ma'lumotni chalkashtirib yubormasligi uchun kerak.",
            "Bu parallelizmni cheklash emas, balki parallel ishlashda ma'lumot yaxlitligini saqlashdir.",
        ],
    )
    add_code(
        slide,
        """_MANUAL_CORRECT_LOCKS: Dict[int, asyncio.Lock] = {}

def _manual_correct_lock(user_id: int) -> asyncio.Lock:
    lock = _MANUAL_CORRECT_LOCKS.get(user_id)
    if lock is None:
        lock = asyncio.Lock()
        _MANUAL_CORRECT_LOCKS[user_id] = lock
    return lock""",
        font_size=12.5,
    )
    add_footer(slide, "Funksiya: asyncio.Lock")

    add_section_slide(
        prs,
        "10. Xulosa",
        "Taqdimotda aytish uchun qisqa gaplar",
        [
            "Botda parallelizm asosan asyncio orqali tashkil qilingan.",
            "Eng katta tezlik yutug'i AI batchlarni asyncio.gather bilan parallel yuborishda.",
            "Fon tasklar quiz, keepalive va progress animatsiyani asosiy oqimdan ajratadi.",
            "to_thread bloklovchi fayl ishlovlarini event loopdan olib chiqadi.",
            "Lock va Event parallel muhitda tartib va boshqaruvni saqlaydi.",
        ],
    )

    prs.save(OUT)
    print(OUT.resolve())


if __name__ == "__main__":
    main()
